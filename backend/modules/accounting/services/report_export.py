"""Report export utilities — PDF (ReportLab) and Excel (openpyxl) — Phase 13.

Both functions take the SAME plain-dict/list shape ``FinancialStatementService``/
``ReportService`` methods already return (no separate "export DTOs") and are
purely presentational — read-only, never touch the database, never mutate
financial history (spec.md requirement: "reports do not mutate financial
history").

Spec ref: specs/008-accounting-finance/spec.md §17
Tasks ref: specs/008-accounting-finance/tasks.md T258
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _fmt_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:,.2f}"
    if isinstance(value, date | datetime):
        return value.isoformat()
    return str(value)


def _rows_for_template(
    report_data: dict[str, Any], template: str
) -> tuple[list[str], list[list[Any]]]:
    """Flatten a report dict into (headers, rows) for a known report template.

    Supported templates: ``trial_balance``, ``balance_sheet``, ``profit_loss``,
    ``gl_report``, ``ledger_statement`` (customer/supplier ledger, bank book,
    cash book — all share the same opening/transactions/closing shape),
    ``journal_report``. Unknown templates fall back to a generic key/value
    dump so export never hard-fails on a report shape it doesn't specially know.
    """
    rows: list[list[Any]]
    if template == "trial_balance":
        headers = ["Account Code", "Total Debit", "Total Credit"]
        rows = [
            [r["account_code"], r["total_debit"], r["total_credit"]]
            for r in report_data["rows"]
        ]
        return headers, rows

    if template == "balance_sheet":
        headers = ["Section", "Account Code", "Account Name", "Amount"]
        rows = []
        for section, label in (
            ("assets", "Assets"),
            ("liabilities", "Liabilities"),
            ("equity", "Equity"),
        ):
            for r in report_data[section]:
                rows.append([label, r["account_code"], r["account_name"], r["amount"]])
        return headers, rows

    if template == "profit_loss":
        headers = ["Section", "Account Code", "Account Name", "Amount"]
        rows = []
        for section, label in (
            ("revenue_lines", "Revenue"),
            ("expense_lines", "Expense"),
        ):
            for r in report_data[section]:
                rows.append([label, r["account_code"], r["account_name"], r["amount"]])
        return headers, rows

    if template == "gl_report":
        headers = [
            "Posting Date",
            "Journal #",
            "Account Code",
            "Debit",
            "Credit",
            "Description",
        ]
        rows = [
            [
                r["posting_date"],
                r.get("journal_number"),
                r["account_code"],
                r["debit_amount"],
                r["credit_amount"],
                r.get("description"),
            ]
            for r in report_data.get("items", report_data.get("rows", []))
        ]
        return headers, rows

    if template == "ledger_statement":
        headers = ["Date", "Reference", "Amount", "Outstanding/Balance"]
        rows = [
            [
                getattr(t, "transaction_date", None),
                getattr(t, "invoice_number", None)
                or getattr(t, "bill_number", None)
                or getattr(t, "reference", None)
                or getattr(t, "transaction_type", None),
                getattr(t, "amount_base", None) or getattr(t, "amount", None),
                getattr(t, "outstanding_amount", None),
            ]
            for t in report_data.get("transactions", [])
        ]
        rows.insert(0, ["", "Opening Balance", report_data.get("opening_balance"), ""])
        rows.append(["", "Closing Balance", report_data.get("closing_balance"), ""])
        return headers, rows

    if template == "journal_report":
        headers = [
            "Journal #",
            "Posting Date",
            "Type",
            "Status",
            "Debit",
            "Credit",
            "Description",
        ]
        rows = [
            [
                getattr(e, "journal_number", None),
                getattr(e, "posting_date", None),
                getattr(e, "journal_type", None),
                getattr(e, "status", None),
                getattr(e, "total_debit_base", None),
                getattr(e, "total_credit_base", None),
                getattr(e, "description", None),
            ]
            for e in report_data.get("entries", [])
        ]
        return headers, rows

    if template == "audit_log":
        headers = [
            "Occurred At",
            "Entity Type",
            "Entity ID",
            "Action",
            "Actor User ID",
            "Reason",
        ]
        rows = [
            [
                getattr(e, "occurred_at", None),
                getattr(e, "entity_type", None),
                str(getattr(e, "entity_id", "") or ""),
                getattr(e, "action", None),
                str(getattr(e, "actor_user_id", "") or ""),
                getattr(e, "reason", None),
            ]
            for e in report_data.get("entries", [])
        ]
        return headers, rows

    headers = ["Field", "Value"]
    rows = [[k, v] for k, v in report_data.items() if not isinstance(v, list | dict)]
    return headers, rows


def export_to_excel(report_data: dict[str, Any], template: str = "generic") -> bytes:
    """Render ``report_data`` as an .xlsx workbook and return its bytes."""
    headers, rows = _rows_for_template(report_data, template)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = template[:31] or "Report"

    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            if isinstance(value, Decimal):
                sheet.cell(row=row_idx, column=col_idx, value=float(value))
            elif isinstance(value, date | datetime):
                sheet.cell(row=row_idx, column=col_idx, value=value.isoformat())
            else:
                sheet.cell(row=row_idx, column=col_idx, value=value)

    for col_idx in range(1, len(headers) + 1):
        sheet.column_dimensions[
            sheet.cell(row=1, column=col_idx).column_letter
        ].width = 20

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_to_pdf(report_data: dict[str, Any], template: str = "generic") -> bytes:
    """Render ``report_data`` as a PDF document and return its bytes."""
    headers, rows = _rows_for_template(report_data, template)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    elements: list[Any] = [
        Paragraph(template.replace("_", " ").title(), styles["Title"]),
        Spacer(1, 0.25 * inch),
    ]

    table_data = [headers] + [[_fmt_cell(cell) for cell in row] for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f3f4f6")],
                ),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buffer.getvalue()
