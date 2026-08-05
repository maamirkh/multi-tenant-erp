"""Delivery Note export service — Phase 9.

Exports delivery notes as structured text or Excel.

Gated by feature flag: sales.dn_pdf_export

Spec ref: specs/007-sales-management/spec.md §38 Import & Export
Task: T228
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data transfer objects
# ---------------------------------------------------------------------------


@dataclass
class DeliveryLineItem:
    """A single line on the delivery note."""

    line_number: int
    product_code: str
    description: str
    quantity_ordered: Decimal
    quantity_delivered: Decimal
    unit: str = "EA"
    lot_number: str = ""
    serial_numbers: str = ""


@dataclass
class DeliveryExportData:
    """Complete data bundle for delivery note export."""

    delivery_number: str
    delivery_date: str
    dispatch_date: str = ""
    status: str = ""
    order_number: str = ""
    customer_code: str = ""
    customer_name: str = ""
    customer_address: str = ""
    shipping_address: str = ""
    carrier: str = ""
    tracking_number: str = ""
    notes: str = ""
    company_name: str = ""
    company_address: str = ""
    lines: list[DeliveryLineItem] = field(default_factory=list)


# ---------------------------------------------------------------------------
# DeliveryExportService
# ---------------------------------------------------------------------------


class DeliveryExportService:
    """Export delivery notes to various formats.

    Feature flag ``sales.dn_pdf_export`` must be enabled at the API layer
    before calling this service.
    """

    def build_export_data(
        self,
        delivery_note: Any,
        lines: list[Any],
        customer_name: str = "",
        customer_address: str = "",
        shipping_address: str = "",
        company_name: str = "",
        company_address: str = "",
    ) -> DeliveryExportData:
        """Build DeliveryExportData from ORM objects."""
        line_items = []
        for i, line in enumerate(lines, start=1):
            line_items.append(
                DeliveryLineItem(
                    line_number=i,
                    product_code=getattr(line, "product_code", "") or "",
                    description=getattr(line, "description", "") or "",
                    quantity_ordered=Decimal(str(getattr(line, "quantity_ordered", 0))),
                    quantity_delivered=Decimal(
                        str(getattr(line, "quantity_delivered", 0))
                    ),
                    unit=getattr(line, "unit", "EA") or "EA",
                    lot_number=getattr(line, "lot_number", "") or "",
                    serial_numbers=getattr(line, "serial_numbers", "") or "",
                )
            )

        return DeliveryExportData(
            delivery_number=delivery_note.delivery_number,
            delivery_date=str(getattr(delivery_note, "delivery_date", "") or ""),
            dispatch_date=str(getattr(delivery_note, "dispatch_date", "") or ""),
            status=delivery_note.status,
            order_number=str(getattr(delivery_note, "order_id", "") or ""),
            customer_code=str(getattr(delivery_note, "customer_id", "") or ""),
            customer_name=customer_name,
            customer_address=customer_address,
            shipping_address=shipping_address,
            carrier=getattr(delivery_note, "carrier", "") or "",
            tracking_number=getattr(delivery_note, "tracking_number", "") or "",
            notes=getattr(delivery_note, "notes", "") or "",
            company_name=company_name,
            company_address=company_address,
            lines=line_items,
        )

    # ------------------------------------------------------------------
    # Export formats
    # ------------------------------------------------------------------

    def export_text(self, data: DeliveryExportData) -> tuple[bytes, str, str]:
        """Export delivery note as formatted plain text.

        Returns (file_bytes, filename, content_type).
        """
        content = self._render_text(data)
        filename = f"delivery_note_{data.delivery_number}.txt"
        return content.encode("utf-8"), filename, "text/plain; charset=utf-8"

    def export_excel(self, data: DeliveryExportData) -> tuple[bytes, str, str]:
        """Export delivery note as Excel workbook.

        Returns (file_bytes, filename, content_type).
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is required for Excel export. "
                "Install it with: pip install openpyxl"
            ) from exc

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Delivery Note"

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        bold = Font(bold=True)

        # Company header
        ws.merge_cells("A1:G1")
        ws["A1"] = data.company_name or "DevSphere ERP"
        ws["A1"].font = Font(bold=True, size=14)

        if data.company_address:
            ws.merge_cells("A2:G2")
            ws["A2"] = data.company_address

        ws.merge_cells("A4:G4")
        ws["A4"] = f"DELIVERY NOTE: {data.delivery_number}"
        ws["A4"].font = Font(bold=True, size=13)

        meta = [
            ("Delivery Date:", data.delivery_date),
            ("Dispatch Date:", data.dispatch_date),
            ("Status:", data.status),
            ("Sales Order:", data.order_number),
        ]
        row = 6
        for label, value in meta:
            ws.cell(row=row, column=1, value=label).font = bold
            ws.cell(row=row, column=2, value=value)
            row += 1

        if data.carrier:
            ws.cell(row=row, column=1, value="Carrier:").font = bold
            ws.cell(row=row, column=2, value=data.carrier)
            row += 1
        if data.tracking_number:
            ws.cell(row=row, column=1, value="Tracking #:").font = bold
            ws.cell(row=row, column=2, value=data.tracking_number)
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Deliver To:").font = bold
        ws.cell(row=row, column=2, value=data.customer_name or data.customer_code)
        row += 1
        if data.shipping_address:
            for addr_line in data.shipping_address.split("\n"):
                ws.cell(row=row, column=2, value=addr_line)
                row += 1

        # Lines header
        row += 1
        line_headers = [
            "#",
            "Product",
            "Description",
            "Unit",
            "Ordered",
            "Delivered",
            "Lot #",
            "Serial #",
        ]
        for col, hdr in enumerate(line_headers, start=1):
            cell = ws.cell(row=row, column=col, value=hdr)
            cell.fill = header_fill
            cell.font = header_font
        row += 1

        for line in data.lines:
            ws.cell(row=row, column=1, value=line.line_number)
            ws.cell(row=row, column=2, value=line.product_code)
            ws.cell(row=row, column=3, value=line.description)
            ws.cell(row=row, column=4, value=line.unit)
            ws.cell(row=row, column=5, value=float(line.quantity_ordered))
            ws.cell(row=row, column=6, value=float(line.quantity_delivered))
            ws.cell(row=row, column=7, value=line.lot_number)
            ws.cell(row=row, column=8, value=line.serial_numbers)
            row += 1

        if data.notes:
            row += 1
            ws.cell(row=row, column=1, value="Notes:").font = bold
            ws.merge_cells(f"B{row}:H{row}")
            ws.cell(row=row, column=2, value=data.notes)

        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) for cell in col if cell.value is not None),
                default=10,
            )
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"delivery_note_{data.delivery_number}.xlsx"
        return (
            buf.read(),
            filename,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # ------------------------------------------------------------------
    # Text renderer
    # ------------------------------------------------------------------

    def _render_text(self, data: DeliveryExportData) -> str:
        lines: list[str] = []
        width = 72

        def sep(char: str = "=") -> str:
            return char * width

        def row(label: str, value: str, w: int = 28) -> str:
            return f"{label:<{w}}{value}"

        lines.append(sep())
        if data.company_name:
            lines.append(data.company_name.center(width))
        if data.company_address:
            lines.append(data.company_address.center(width))
        lines.append(sep())
        lines.append("DELIVERY NOTE".center(width))
        lines.append(sep())
        lines.append("")

        lines.append(row("Delivery Number:", data.delivery_number))
        lines.append(row("Delivery Date:", data.delivery_date))
        if data.dispatch_date:
            lines.append(row("Dispatch Date:", data.dispatch_date))
        lines.append(row("Status:", data.status))
        if data.order_number:
            lines.append(row("Sales Order #:", data.order_number))
        if data.carrier:
            lines.append(row("Carrier:", data.carrier))
        if data.tracking_number:
            lines.append(row("Tracking #:", data.tracking_number))

        lines.append("")
        lines.append(sep("-"))
        lines.append("DELIVER TO:")
        if data.customer_name:
            lines.append(f"  {data.customer_name}")
        if data.customer_code:
            lines.append(f"  Code: {data.customer_code}")
        address = data.shipping_address or data.customer_address
        if address:
            for addr_line in address.split("\n"):
                lines.append(f"  {addr_line}")

        if data.lines:
            lines.append("")
            lines.append(sep("-"))
            lines.append(
                f"{'#':<4}{'Product':<16}{'Description':<24}"
                f"{'Unit':>6}{'Ordered':>10}{'Delivered':>10}{'Lot':>8}"
            )
            lines.append(sep("-"))
            for line in data.lines:
                lines.append(
                    f"{line.line_number:<4}{line.product_code:<16}"
                    f"{line.description[:23]:<24}"
                    f"{line.unit:>6}"
                    f"{float(line.quantity_ordered):>10.2f}"
                    f"{float(line.quantity_delivered):>10.2f}"
                    f"{line.lot_number:>8}"
                )
                if line.serial_numbers:
                    lines.append(f"{'':>60}S/N: {line.serial_numbers}")

        if data.notes:
            lines.append("")
            lines.append(sep("-"))
            lines.append("Notes:")
            lines.append(data.notes)

        lines.append("")
        lines.append(sep("="))
        lines.append("Received in good condition:".center(width))
        lines.append("")
        lines.append(f"{'Signature: _____________________':>72}")
        lines.append(f"{'Date: ___________________________':>72}")
        lines.append(sep("="))

        return "\n".join(lines)
