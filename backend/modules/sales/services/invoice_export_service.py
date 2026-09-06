"""Invoice export service — Phase 9.

Provides structured invoice export with company branding, line items,
totals, amount-in-words, and tax breakdown.

Export formats:
  - text/plain: well-structured text invoice (default; no external deps)
  - application/vnd.openxmlformats-officedocument.spreadsheetml.sheet: Excel

Production PDF generation (weasyprint/reportlab) can be substituted by
converting the text export through a headless browser or PDF printer
without modifying this service's interface.

Feature flag: sales.invoice_pdf_export must be enabled.

Spec ref: specs/007-sales-management/spec.md §38 Import & Export
Task: T227
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Data transfer objects (service-layer, not ORM)
# ---------------------------------------------------------------------------


@dataclass
class CompanyBranding:
    """Company information used to brand the invoice header."""

    name: str
    address: str = ""
    phone: str = ""
    email: str = ""
    tax_number: str = ""
    logo_text: str = ""  # ASCII branding if no logo image


@dataclass
class InvoiceLineItem:
    """Single line on the exported invoice."""

    line_number: int
    product_code: str
    description: str
    quantity: Decimal
    unit_price: Decimal
    discount_amount: Decimal = field(default_factory=lambda: Decimal("0"))
    tax_amount: Decimal = field(default_factory=lambda: Decimal("0"))
    line_total: Decimal = field(default_factory=lambda: Decimal("0"))
    unit: str = "EA"


@dataclass
class InvoiceExportData:
    """Complete data bundle passed to the export service."""

    invoice_number: str
    invoice_date: str
    due_date: str
    status: str
    currency_code: str
    customer_code: str
    customer_name: str
    customer_address: str = ""
    payment_term: str = ""
    reference: str = ""
    subtotal: Decimal = field(default_factory=lambda: Decimal("0"))
    discount_amount: Decimal = field(default_factory=lambda: Decimal("0"))
    tax_amount: Decimal = field(default_factory=lambda: Decimal("0"))
    charges_amount: Decimal = field(default_factory=lambda: Decimal("0"))
    total_amount: Decimal = field(default_factory=lambda: Decimal("0"))
    amount_in_words: str = ""
    notes: str = ""
    lines: list[InvoiceLineItem] = field(default_factory=list)
    company: CompanyBranding = field(default_factory=lambda: CompanyBranding(name=""))


# ---------------------------------------------------------------------------
# InvoiceExportService
# ---------------------------------------------------------------------------


class InvoiceExportService:
    """Export sales invoices to various formats.

    Usage::

        svc = InvoiceExportService()
        data = InvoiceExportData(...)
        file_bytes, filename, content_type = svc.export_text(data)

        # Or build data from ORM models:
        data = svc.build_export_data(invoice, lines, company_branding)
        file_bytes, filename, content_type = svc.export_text(data)
    """

    # ------------------------------------------------------------------
    # Public factory helpers
    # ------------------------------------------------------------------

    def build_export_data(
        self,
        invoice: Any,
        lines: list[Any],
        company: CompanyBranding | None = None,
        customer_name: str = "",
        customer_address: str = "",
        payment_term: str = "",
    ) -> InvoiceExportData:
        """Build an InvoiceExportData from ORM invoice + line objects."""
        line_items = []
        for i, line in enumerate(lines, start=1):
            line_items.append(
                InvoiceLineItem(
                    line_number=i,
                    product_code=getattr(line, "product_code", "") or "",
                    description=getattr(line, "description", "") or "",
                    quantity=Decimal(str(getattr(line, "quantity_invoiced", 1))),
                    unit_price=Decimal(str(getattr(line, "unit_price", 0))),
                    discount_amount=Decimal(str(getattr(line, "discount_amount", 0))),
                    tax_amount=Decimal(str(getattr(line, "tax_amount", 0))),
                    line_total=Decimal(str(getattr(line, "line_total", 0))),
                    unit=getattr(line, "unit", "EA") or "EA",
                )
            )

        return InvoiceExportData(
            invoice_number=invoice.invoice_number,
            invoice_date=str(invoice.invoice_date),
            due_date=str(invoice.due_date),
            status=invoice.status,
            currency_code=invoice.currency_code,
            customer_code=str(getattr(invoice, "customer_id", "")),
            customer_name=customer_name,
            customer_address=customer_address,
            payment_term=payment_term,
            reference=getattr(invoice, "reference", "") or "",
            subtotal=Decimal(str(invoice.subtotal)),
            discount_amount=Decimal(str(invoice.discount_amount)),
            tax_amount=Decimal(str(invoice.tax_amount)),
            charges_amount=Decimal(str(getattr(invoice, "charges_amount", 0) or 0)),
            total_amount=Decimal(str(invoice.total_amount)),
            amount_in_words=getattr(invoice, "amount_in_words", "") or "",
            notes=getattr(invoice, "notes", "") or "",
            lines=line_items,
            company=company or CompanyBranding(name=""),
        )

    # ------------------------------------------------------------------
    # Export formats
    # ------------------------------------------------------------------

    def export_text(self, data: InvoiceExportData) -> tuple[bytes, str, str]:
        """Export invoice as formatted plain text.

        Returns (file_bytes, filename, content_type).
        """
        content = self._render_text(data)
        filename = f"invoice_{data.invoice_number}.txt"
        return content.encode("utf-8"), filename, "text/plain; charset=utf-8"

    def export_excel(self, data: InvoiceExportData) -> tuple[bytes, str, str]:
        """Export invoice as Excel workbook using openpyxl.

        Returns (file_bytes, filename, content_type).
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is required for Excel export. "
                "Install it with: pip install openpyxl"
            ) from exc

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Invoice"

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        section_font = Font(bold=True, size=11)

        # --- Company header ---
        ws.merge_cells("A1:G1")
        ws["A1"] = data.company.name or "DevSphere ERP"
        ws["A1"].font = Font(bold=True, size=14)
        if data.company.address:
            ws.merge_cells("A2:G2")
            ws["A2"] = data.company.address

        # --- Invoice title ---
        ws.merge_cells("A4:G4")
        ws["A4"] = f"INVOICE: {data.invoice_number}"
        ws["A4"].font = Font(bold=True, size=13)

        # --- Invoice meta ---
        meta = [
            ("Invoice Date:", data.invoice_date),
            ("Due Date:", data.due_date),
            ("Status:", data.status),
            ("Currency:", data.currency_code),
        ]
        row = 6
        for label, value in meta:
            ws.cell(row=row, column=1, value=label).font = section_font
            ws.cell(row=row, column=2, value=value)
            row += 1

        # --- Customer info ---
        row += 1
        ws.cell(row=row, column=1, value="Bill To:").font = section_font
        ws.cell(row=row, column=2, value=data.customer_name or data.customer_code)
        row += 1
        if data.customer_address:
            ws.cell(row=row, column=2, value=data.customer_address)
            row += 1

        # --- Lines header ---
        row += 1
        line_headers = [
            "#",
            "Product",
            "Description",
            "Qty",
            "Unit",
            "Unit Price",
            "Discount",
            "Tax",
            "Total",
        ]
        for col, hdr in enumerate(line_headers, start=1):
            cell = ws.cell(row=row, column=col, value=hdr)
            cell.fill = header_fill
            cell.font = header_font

        # --- Lines ---
        row += 1
        for line in data.lines:
            ws.cell(row=row, column=1, value=line.line_number)
            ws.cell(row=row, column=2, value=line.product_code)
            ws.cell(row=row, column=3, value=line.description)
            ws.cell(row=row, column=4, value=float(line.quantity))
            ws.cell(row=row, column=5, value=line.unit)
            ws.cell(row=row, column=6, value=float(line.unit_price))
            ws.cell(row=row, column=7, value=float(line.discount_amount))
            ws.cell(row=row, column=8, value=float(line.tax_amount))
            ws.cell(row=row, column=9, value=float(line.line_total))
            row += 1

        # --- Totals ---
        row += 1
        totals = [
            ("Subtotal:", float(data.subtotal)),
            ("Discount:", -float(data.discount_amount)),
            ("Tax:", float(data.tax_amount)),
            ("Charges:", float(data.charges_amount)),
            ("TOTAL:", float(data.total_amount)),
        ]
        for label, amount in totals:
            ws.cell(row=row, column=8, value=label).font = section_font
            ws.cell(row=row, column=9, value=amount)
            row += 1

        # --- Amount in words ---
        if data.amount_in_words:
            row += 1
            ws.cell(row=row, column=1, value="Amount in Words:").font = section_font
            ws.merge_cells(f"B{row}:I{row}")
            ws.cell(row=row, column=2, value=data.amount_in_words)

        # Auto-fit columns
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) for cell in col if cell.value is not None),
                default=10,
            )
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"invoice_{data.invoice_number}.xlsx"
        return (
            buf.read(),
            filename,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # ------------------------------------------------------------------
    # Internal renderers
    # ------------------------------------------------------------------

    def _render_text(self, data: InvoiceExportData) -> str:
        lines: list[str] = []
        width = 72

        def sep(char: str = "=") -> str:
            return char * width

        def row(label: str, value: str, width: int = 30) -> str:
            return f"{label:<{width}}{value}"

        lines.append(sep())

        # Company branding
        if data.company.name:
            lines.append(data.company.name.center(width))
        if data.company.address:
            lines.append(data.company.address.center(width))
        if data.company.phone:
            lines.append(f"Tel: {data.company.phone}".center(width))
        if data.company.email:
            lines.append(f"Email: {data.company.email}".center(width))
        if data.company.tax_number:
            lines.append(f"Tax Reg: {data.company.tax_number}".center(width))

        lines.append(sep())
        lines.append("INVOICE".center(width))
        lines.append(sep())
        lines.append("")

        # Invoice metadata
        lines.append(row("Invoice Number:", data.invoice_number))
        lines.append(row("Invoice Date:", data.invoice_date))
        lines.append(row("Due Date:", data.due_date))
        lines.append(row("Status:", data.status))
        lines.append(row("Currency:", data.currency_code))
        if data.reference:
            lines.append(row("Reference:", data.reference))
        if data.payment_term:
            lines.append(row("Payment Terms:", data.payment_term))

        lines.append("")
        lines.append(sep("-"))
        lines.append("BILL TO:")
        if data.customer_name:
            lines.append(f"  {data.customer_name}")
        if data.customer_code:
            lines.append(f"  Code: {data.customer_code}")
        if data.customer_address:
            for addr_line in data.customer_address.split("\n"):
                lines.append(f"  {addr_line}")

        # Line items
        if data.lines:
            lines.append("")
            lines.append(sep("-"))
            lines.append(
                f"{'#':<4}{'Product':<16}{'Description':<24}{'Qty':>6}{'Unit':>6}"
                f"{'Price':>10}{'Disc':>8}{'Tax':>8}{'Total':>10}"
            )
            lines.append(sep("-"))
            for line in data.lines:
                lines.append(
                    f"{line.line_number:<4}{line.product_code:<16}"
                    f"{line.description[:23]:<24}"
                    f"{float(line.quantity):>6.2f}{line.unit:>6}"
                    f"{float(line.unit_price):>10.2f}"
                    f"{float(line.discount_amount):>8.2f}"
                    f"{float(line.tax_amount):>8.2f}"
                    f"{float(line.line_total):>10.2f}"
                )

        # Totals
        lines.append("")
        lines.append(sep("-"))
        col_w = 40
        lines.append(
            row("Subtotal:", f"{data.currency_code} {data.subtotal:>12.2f}", col_w)
        )
        if data.discount_amount:
            lines.append(
                row(
                    "Discount:",
                    f"{data.currency_code} -{data.discount_amount:>11.2f}",
                    col_w,
                )
            )
        if data.tax_amount:
            lines.append(
                row("Tax:", f"{data.currency_code} {data.tax_amount:>12.2f}", col_w)
            )
        if data.charges_amount:
            lines.append(
                row(
                    "Charges:",
                    f"{data.currency_code} {data.charges_amount:>12.2f}",
                    col_w,
                )
            )
        lines.append(sep("-"))
        lines.append(
            row(
                "TOTAL AMOUNT:",
                f"{data.currency_code} {data.total_amount:>12.2f}",
                col_w,
            )
        )
        lines.append(sep("="))

        if data.amount_in_words:
            lines.append(f"Amount in Words: {data.amount_in_words}")
            lines.append("")

        if data.notes:
            lines.append("Notes:")
            lines.append(data.notes)
            lines.append("")

        lines.append(sep("="))
        lines.append("Thank you for your business.".center(width))
        lines.append(sep("="))

        return "\n".join(lines)
