"""T200 (FR-RPT-219) — a filter scope matching zero rows still produces a
valid file: the row schema's headers and zero data rows, never an error."""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from modules.reports.schemas.sales import SalesByCustomerRow
from tests.integration.api.v1.reports.conftest import (
    auth_header,
    export_url,
    parse_csv,
    setup_company,
)


def test_empty_csv_export_is_headers_only(
    test_client: TestClient, db_session: Session
) -> None:
    token, company_id = setup_company(db_session, test_client)

    resp = test_client.get(
        export_url(company_id, "sales.by_customer", "csv"), headers=auth_header(token)
    )

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    rows = parse_csv(resp.content)
    assert rows == [list(SalesByCustomerRow.model_fields)]


def test_empty_xlsx_export_is_headers_only(
    test_client: TestClient, db_session: Session
) -> None:
    token, company_id = setup_company(db_session, test_client)

    resp = test_client.get(
        export_url(company_id, "sales.by_customer", "xlsx"), headers=auth_header(token)
    )

    assert resp.status_code == 200, resp.text
    sheet = load_workbook(io.BytesIO(resp.content), read_only=True).active
    assert sheet is not None
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    assert rows == [list(SalesByCustomerRow.model_fields)]
