"""T210 — exporting a Category D (aggregate) report never invokes
``count_export_rows()``/``iter_export_rows()`` (spy-verified): one
``run()``, fixed shape serialized directly. Also covers T194's PDF
delegation and format rejection for aggregates."""

from __future__ import annotations

import io
from collections.abc import Iterator
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy.orm import Session

from modules.reports.services.adapters.accounting_adapter import AccountingAdapter
from modules.reports.services.adapters.inventory_adapter import InventoryAdapter
from tests.integration.api.v1.accounting.test_reports_api import _setup
from tests.integration.api.v1.reports.conftest import (
    auth_header,
    export_url,
    parse_csv,
    setup_company,
)


@pytest.fixture
def seam_calls(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    calls: list[str] = []

    def count_spy(
        self: object, db: Session, company_id: UUID, report_key: str, filters: BaseModel
    ) -> int:
        calls.append(f"count:{report_key}")
        raise AssertionError("aggregate export must not count rows")

    def iter_spy(
        self: object,
        db: Session,
        company_id: UUID,
        report_key: str,
        filters: BaseModel,
        sort: str | None,
        batch_size: int,
    ) -> Iterator[list[BaseModel]]:
        calls.append(f"iter:{report_key}")
        raise AssertionError("aggregate export must not iterate rows")

    for adapter_cls in (InventoryAdapter, AccountingAdapter):
        monkeypatch.setattr(adapter_cls, "count_export_rows", count_spy)
        monkeypatch.setattr(adapter_cls, "iter_export_rows", iter_spy)
    return calls


def test_inventory_valuation_csv_is_single_run(
    test_client: TestClient, db_session: Session, seam_calls: list[str]
) -> None:
    token, company_id = setup_company(db_session, test_client)

    resp = test_client.get(
        export_url(company_id, "inventory.valuation", "csv"), headers=auth_header(token)
    )

    assert resp.status_code == 200, resp.text
    rows = parse_csv(resp.content)
    assert rows[0] == ["field", "value"]
    assert "grand_total_value" in {r[0] for r in rows[1:]}
    assert seam_calls == []


def test_trial_balance_xlsx_and_pdf(
    test_client: TestClient, db_session: Session, seam_calls: list[str]
) -> None:
    token, company_id = setup_company(db_session, test_client)
    fixture = _setup(db_session, company_id)
    scope = f"&filters[period_id]={fixture['period'].id}"
    headers = auth_header(token)

    xlsx = test_client.get(
        export_url(company_id, "accounting.trial_balance", "xlsx") + scope,
        headers=headers,
    )
    assert xlsx.status_code == 200, xlsx.text
    sheet = load_workbook(io.BytesIO(xlsx.content), read_only=True).active
    assert sheet is not None
    first = next(sheet.iter_rows(values_only=True))
    assert list(first) == ["field", "value"]

    pdf = test_client.get(
        export_url(company_id, "accounting.trial_balance", "pdf") + scope,
        headers=headers,
    )
    assert pdf.status_code == 200, pdf.text
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")

    csv_resp = test_client.get(
        export_url(company_id, "accounting.trial_balance", "csv") + scope,
        headers=headers,
    )
    assert csv_resp.status_code == 422, csv_resp.text
    assert csv_resp.json()["error"]["code"] == "REPORT_INVALID_FILTER"
    assert seam_calls == []


def test_non_pdf_report_rejects_pdf_format(
    test_client: TestClient, db_session: Session
) -> None:
    """AR aging advertises CSV/XLSX only (Phase 6 registry correction) — a
    PDF request is a clean 422, never an adapter-level crash."""
    token, company_id = setup_company(db_session, test_client)

    resp = test_client.get(
        export_url(company_id, "accounting.ar_aging", "pdf")
        + "&filters[as_of_date]=2026-01-31",
        headers=auth_header(token),
    )

    assert resp.status_code == 422, resp.text
    assert resp.json()["error"]["code"] == "REPORT_INVALID_FILTER"


def test_report_without_export_permission_code_is_denied(
    test_client: TestClient, db_session: Session
) -> None:
    """``sales.kpis`` has no ``export_permission`` at all — never exportable."""
    token, company_id = setup_company(db_session, test_client)

    resp = test_client.get(
        export_url(company_id, "sales.kpis", "csv"), headers=auth_header(token)
    )

    assert resp.status_code == 403, resp.text
    assert resp.json()["error"]["code"] == "REPORT_PERMISSION_DENIED"
